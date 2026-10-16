import os
from pathlib import Path
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Carga segura de API Key desde variable de entorno o .streamlit/secrets.toml
API_KEY = os.environ.get("GEMINI_API_KEY")
if not API_KEY:
    secrets_path = Path(__file__).parent / ".streamlit" / "secrets.toml"
    if secrets_path.exists():
        try:
            import tomllib
            with open(secrets_path, "rb") as f:
                API_KEY = tomllib.load(f).get("GEMINI_API_KEY")
        except Exception:
            pass

def generar_informe_completo_con_ia():
    # 1. Configurar la API Key
    genai.configure(api_key=API_KEY)

    # Datos completos del paciente de prueba
    paciente = "Pérez, Mateo"
    hc = "402891"
    edad = "9 años, 4 meses"
    escolaridad = "4° Grado Primaria"
    fecha_eval = "Octubre 2026"
    profesional = "Lic. Juan Machado"
    
    # Grilla cuantitativa estructurada (Batería, Subtest/Índice, Puntaje, Z, Rango)
    filas_resultados = [
        {"bateria": "WISC-V", "subtest": "Comprensión Verbal (ICV)", "valor": "106 (CI)", "z": "+0.40", "rango": "Normal"},
        {"bateria": "WISC-V", "subtest": "Visuoespacial (IVE)", "valor": "97 (CI)", "z": "-0.20", "rango": "Normal"},
        {"bateria": "WISC-V", "subtest": "Razonamiento Fluido (IRF)", "valor": "100 (CI)", "z": "0.00", "rango": "Normal"},
        {"bateria": "WISC-V", "subtest": "Memoria de Trabajo (IMT)", "valor": "78 (CI)", "z": "-1.47", "rango": "Dificultad"},
        {"bateria": "WISC-V", "subtest": "Velocidad de Procesamiento (IVP)", "valor": "81 (CI)", "z": "-1.27", "rango": "Dificultad"},
        {"bateria": "CPT-II", "subtest": "Omisiones (Atención Sostenida)", "valor": "66 (T)", "z": "-1.60", "rango": "Dificultad"},
        {"bateria": "CPT-II", "subtest": "Comisiones (Control Inhibitorio)", "valor": "52 (T)", "z": "-0.20", "rango": "Normal"}
    ]
    
    observacion_conducta = "Ingresó al consultorio con adecuada autonomía y buen contacto vincular. Durante la administración de las primeras pruebas mantuvo la atención y el interés. Sin embargo, hacia los 25-30 minutos de evaluación (coincidiendo con las tareas de alta demanda en memoria operativa y velocidad), evidenció signos de fatiga física, inquietud motora leve y verbalizaciones de agobio ('es muy largo'), requiriendo instigación verbal y pausas breves para sostener el set de trabajo."

    # 2. Prompt estructurado y exigente para obtener teoría y profundidad clínica
    prompt_completo = f"""
    Actuá como un neuropsicólogo y psicopedagogo clínico senior (estilo FLENI). Redactá un informe de evaluación neuropsicológica completo y altamente técnico para:
    - Paciente: {paciente}, Edad: {edad}, Escolaridad: {escolaridad}
    - Resultados Cuantitativos Z: {filas_resultados}
    - Observación de conducta: {observacion_conducta}
    
    ESTRUCTURA REQUERIDA Y PAUTAS DE REDACCIÓN:
    1. MOTIVO DE CONSULTA Y ANTECEDENTES: Redactá un párrafo formal contextualizando la derivación escolar/médica por dificultades atencionales y de rendimiento operativo.
    2. OBSERVACIÓN DE LA CONDUCTA: Desarrollá un párrafo clínico formal integrando los datos de fatiga, ritmo de trabajo y tolerancia a la frustración.
    3. ANÁLISIS CUALITATIVO POR DOMINIOS (Con sustento teórico profundo):
       Para cada uno de los siguientes dominios (Funcionamiento Intelectual, Atención y Funciones Ejecutivas, Memoria y Aprendizaje, Procesamiento Visomotor), escribí primero una introducción teórica conceptual detallada sobre qué evalúa el dominio y qué redes neuroanatómicas o cognitivas implican. Luego, redactá el análisis específico del paciente usando **negritas** únicamente en los calificadores de desempeño (ej. **adecuado**, **dificultades**, **por debajo**, **acorde**). Explicá de forma rigurosa el impacto de la marcada dispersión observada entre los índices operativos (IMT, IVP) y los cristalizados.
    4. CONCLUSIONES NEUROPSICOLÓGICAS: Integración formal del perfil intrapersonal (puntos fuertes y descensos relativos). Recordá cerrar indicando que el diagnóstico clínico definitivo queda a cargo del médico tratante.
    5. SUGERENCIAS Y RECOMENDACIONES: Detallá pautas terapéuticas concretas (psicopedagogía, psicoterapia) y adaptaciones metodológicas y de acceso específicas para el ámbito escolar.
    """

    print("🧠 Consultando a Gemini con enfoque teórico profundo...")
    model = genai.GenerativeModel('gemini-3.8-flash')
    response = model.generate_content(prompt_completo)
    texto_ia = response.text

    # 3. Creación del documento Word con diseño institucional profesional
    doc = Document()
    
    # Márgenes estilo institucional (1 pulgada / 2.54 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilo base Arial 10.5
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(30, 30, 30)

    # Encabezado institucional simulado
    p_enc = doc.add_paragraph()
    p_enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_enc = p_enc.add_run("SERVICIO DE NEUROPSICOLOGÍA CLÍNICA\nProtocolo de Evaluación Integral")
    run_enc.font.size = Pt(8.5)
    run_enc.font.color.rgb = RGBColor(100, 100, 100)

    # Título Principal
    titulo = doc.add_heading("INFORME DE EVALUACIÓN NEUROPSICOLÓGICA Y PSICOPEDAGÓGICA", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.runs[0].font.size = Pt(14)
    titulo.runs[0].font.bold = True

    # Sección 1: Datos de Filiación en Tabla Prolija
    doc.add_heading("1. Datos de Filiación", level=2)
    tabla_fil = doc.add_table(rows=4, cols=2)
    tabla_fil.alignment = WD_TABLE_ALIGNMENT.CENTER
    datos_fil = [
        (f"Apellido y Nombre: {paciente}", f"N° HC: {hc}"),
        (f"Edad: {edad}", f"Escolaridad: {escolaridad}"),
        (f"Fecha de Evaluación: {fecha_eval}", f"Solicitante: Equipo Médico / Escolar"),
        (f"Profesional Actuante: {profesional}", "Matrícula Provincial / Nacional")
    ]
    for idx, (c1, c2) in enumerate(datos_fil):
        row = tabla_fil.rows[idx]
        row.cells[0].text = c1
        row.cells[1].text = c2

    doc.add_paragraph() # Espacio

    # Sección 2: Grilla Cuantitativa de Resultados (Tabla Z con diseño)
    doc.add_heading("2. Resultados Cuantitativos (Grilla Z)", level=2)
    tabla_res = doc.add_table(rows=1, cols=5)
    tabla_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Batería / Test", "Subtest / Índice", "Puntaje Bruto / Estándar", "Puntaje Z", "Rango Clínico"]
    hdr_cells = tabla_res.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9.5)

    for item in filas_resultados:
        row_cells = tabla_res.add_row().cells
        row_cells[0].text = item["bateria"]
        row_cells[1].text = item["subtest"]
        row_cells[2].text = item["valor"]
        row_cells[3].text = item["z"]
        row_cells[4].text = item["rango"]
        for cell in row_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)

    doc.add_paragraph() # Espacio

    # Sección 3: Análisis Clínico y Teórico Producido por la IA
    doc.add_heading("3. Análisis Cualitativo, Teórico y Conclusiones", level=2)
    doc.add_paragraph(texto_ia)

    # Firma al pie
    doc.add_paragraph("\n\n________________________________________")
    p_firma = doc.add_paragraph(f"{profesional}\nPsicólogo / Neuropsicólogo Clínico")
    p_firma.runs[0].font.bold = True

    # Guardar archivo definitivo
    nombre_salida = "Informe_Neuropsicologico_Completo.docx"
    doc.save(nombre_salida)
    print(f"¡Éxito total! Informe robusto y completo guardado como: {nombre_salida}")

if __name__ == "__main__":
    generar_informe_completo_con_ia()