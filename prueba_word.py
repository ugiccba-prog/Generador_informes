import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generar_primer_informe():
    doc = Document()
    
    # Configurar fuente base
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    # Encabezado
    titulo = doc.add_heading("INFORME DE PRUEBA INSTITUCIONAL", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Datos de Filiación
    doc.add_heading("1. Datos de Filiación", level=2)
    tabla_fil = doc.add_table(rows=2, cols=2)
    tabla_fil.style = "Table Grid"
    tabla_fil.rows[0].cells[0].text = "Paciente: Pérez, Mateo"
    tabla_fil.rows[0].cells[1].text = "Edad: 9 años, 4 meses"
    tabla_fil.rows[1].cells[0].text = "Evaluador: Lic. Juan Machado"
    tabla_fil.rows[1].cells[1].text = "Fecha: Octubre 2026"

    # Resultados Cuantitativos (Tabla Z)
    doc.add_heading("2. Resultados Cuantitativos (Grilla Z)", level=2)
    tabla_res = doc.add_table(rows=1, cols=4)
    tabla_res.style = "Table Grid"
    
    encabezados = ["Batería / Índice", "Subtest", "Puntaje Z", "Rango Clínico"]
    for i, texto in enumerate(encabezados):
        tabla_res.rows[0].cells[i].paragraphs[0].add_run(texto).bold = True

    # Fila de ejemplo con cálculo Z simulado
    fila_nueva = tabla_res.add_row().cells
    fila_nueva[0].text = "WISC-V"
    fila_nueva[1].text = "Memoria de Trabajo (IMT)"
    fila_nueva[2].text = "-1.47"
    fila_nueva[3].text = "Dificultad"

    # Conclusiones
    doc.add_heading("3. Conclusiones Preliminares", level=2)
    doc.add_paragraph("Este documento fue generado de forma automática y robusta directamente desde un script de Python en tu computadora, respetando los estándares formales exigidos.")

    # Guardar archivo
    nombre_archivo = "Informe_Prueba_Mateo.docx"
    doc.save(nombre_archivo)
    print(f"¡Éxito! Archivo Word generado correctamente como: {nombre_archivo}")

if __name__ == "__main__":
    generar_primer_informe()